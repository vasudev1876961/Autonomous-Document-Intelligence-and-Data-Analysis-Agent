"""
Excel Parser using openpyxl and pandas for structured workbook inspection with sheet and cell provenance.
"""
from pathlib import Path
import uuid
import openpyxl
import pandas as pd
from src.ingestion.metadata import DocumentMetadata, ExtractedChunk, FileType, SourceLocation, compute_file_sha256

class ExcelParser:
    def parse(self, file_path: Path) -> tuple[DocumentMetadata, list[ExtractedChunk]]:
        path_obj = Path(file_path)
        file_hash = compute_file_sha256(path_obj)
        doc_id = f"doc_{file_hash[:12]}"
        chunks: list[ExtractedChunk] = []

        try:
            wb = openpyxl.load_workbook(str(path_obj), data_only=True)
            sheet_names = wb.sheetnames
            
            for sheet_name in sheet_names:
                df = pd.read_excel(str(path_obj), sheet_name=sheet_name)
                if df.empty:
                    continue
                
                # Full sheet representation
                csv_repr = df.to_csv(index=False)
                row_count, col_count = df.shape
                cell_range = f"A1:{openpyxl.utils.get_column_letter(col_count)}{row_count + 1}"
                
                chunks.append(
                    ExtractedChunk(
                        chunk_id=f"{doc_id}_{sheet_name}_full",
                        document_id=doc_id,
                        file_name=path_obj.name,
                        content=f"Workbook Sheet '{sheet_name}' (Range {cell_range}):\n{csv_repr}",
                        content_type="table",
                        location=SourceLocation(
                            document_id=doc_id,
                            file_name=path_obj.name,
                            file_type=FileType.XLSX,
                            sheet_name=sheet_name,
                            cell_range=cell_range
                        ),
                        metadata={
                            "sheet_name": sheet_name,
                            "columns": [str(c) for c in df.columns],
                            "row_count": row_count,
                            "col_count": col_count
                        }
                    )
                )

                # Individual row key-values for granular retrieval
                for r_idx, row in df.iterrows():
                    row_dict = {str(k): (None if pd.isna(v) else v) for k, v in row.to_dict().items()}
                    row_desc = ", ".join([f"{k}: {v}" for k, v in row_dict.items() if v is not None])
                    chunks.append(
                        ExtractedChunk(
                            chunk_id=f"{doc_id}_{sheet_name}_r{r_idx + 2}",
                            document_id=doc_id,
                            file_name=path_obj.name,
                            content=f"[{sheet_name} Row {r_idx + 2}] {row_desc}",
                            content_type="key_value",
                            location=SourceLocation(
                                document_id=doc_id,
                                file_name=path_obj.name,
                                file_type=FileType.XLSX,
                                sheet_name=sheet_name,
                                row_index=r_idx + 2,
                                cell_range=f"Row {r_idx + 2}"
                            ),
                            metadata={"sheet_name": sheet_name, "row_data": row_dict}
                        )
                    )

            wb.close()
            meta = DocumentMetadata(
                document_id=doc_id,
                file_name=path_obj.name,
                file_path=str(path_obj),
                file_type=FileType.XLSX,
                file_size_bytes=path_obj.stat().st_size,
                content_hash=file_hash,
                sheet_names=sheet_names,
                extraction_status="success",
            )
            return meta, chunks

        except Exception as e:
            meta = DocumentMetadata(
                document_id=doc_id,
                file_name=path_obj.name,
                file_path=str(path_obj),
                file_type=FileType.XLSX,
                file_size_bytes=path_obj.stat().st_size if path_obj.exists() else 0,
                content_hash=file_hash,
                extraction_status="failed",
                error_message=str(e),
            )
            return meta, []
